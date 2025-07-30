import os
import numpy as np
import sounddevice as sd
import soundfile as sf
from fastapi import FastAPI, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from datetime import datetime
from docx import Document
from dotenv import load_dotenv
from threading import Thread
from pydantic import BaseModel
from typing import Optional, List

from selenium import webdriver
from selenium.webdriver.chrome.service import Service
from selenium.webdriver.chrome.options import Options
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
import time

from assemblyai_client.client import AssemblyAIClient
from utils.formatter import format_transcript_to_txt
from llm.groq_llm import call_groq_llm
from db.store import store_summary, fetch_all_summaries, get_project_by_id, get_freelancer_by_id, get_meeting_summaries_by_user
from db_connection import DatabaseConnection
from github.uploader import AssignFreelancerPayload, ProjectPayload, MeetingPayload, load_gist, update_gist

# Load environment variables
load_dotenv()
ASSEMBLY_API_KEY = os.getenv("ASSEMBLYAI_API_KEY")

# Paths
AUDIO_PATH = "audio/recorded.wav"
TXT_PATH = "docs/transcript.txt"
DOCX_PATH = "docs/transcript.docx"

# Create necessary directories
os.makedirs("audio", exist_ok=True)
os.makedirs("docs", exist_ok=True)

# FastAPI setup
app = FastAPI(title="Meeting Bot API", version="1.0.0")
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
    expose_headers=["*"],
    allow_credentials=True
)

# Pydantic models
class MeetingStartRequest(BaseModel):
    project_id: int
    freelancer_id: int

class MeetingSummaryRequest(BaseModel):
    freelancer_id: int
    project_id: int
    freelancer_name: str
    project_name: str
    summary: str
    blocker: Optional[str] = None

# Recording state
samplerate = 16000
channels = 1
recording = []
stream = None
is_recording = False
current_meeting_context = None

# AssemblyAI client
client = AssemblyAIClient(ASSEMBLY_API_KEY)

# Initialize database table
def initialize_database():
    try:
        db_connection = DatabaseConnection()
        db_connection.create_meeting_summaries_table()
        print("✅ Database initialized successfully")
    except Exception as e:
        print(f"❌ Database initialization failed: {e}")

# Audio callback
def audio_callback(indata, frames, time, status):
    recording.append(indata.copy())

# Start recording
def start_recording():
    global recording, stream, is_recording
    recording = []
    stream = sd.InputStream(samplerate=samplerate, channels=channels, callback=audio_callback)
    stream.start()
    is_recording = True

# Stop recording and transcribe
def stop_and_transcribe():
    global stream, is_recording, current_meeting_context

    if stream:
        stream.stop()
        stream.close()
        stream = None
    is_recording = False

    audio_np = np.concatenate(recording, axis=0)
    sf.write(AUDIO_PATH, audio_np, samplerate)

    with open(AUDIO_PATH, "rb") as f:
        upload_url = client.upload_file(f)
    resp = client.transcribe_audio(upload_url, diarization=True)
    result = client.poll_transcription(resp["id"])

    # Handle failed transcript
    if not result or result.get("status") != "completed":
        print("❌ Transcription failed:", result)
        return {"status": "failed", "error": result.get("error", "Unknown transcription error")}

    # Format text from result
    raw_txt = format_transcript_to_txt(result)

    with open(TXT_PATH, "w", encoding="utf-8") as f:
        f.write(raw_txt)

    os.remove(AUDIO_PATH)

    # LLM Summary
    llm_res = call_groq_llm(raw_txt)
    if llm_res.get("status") == "success":
        # Store in database with meeting context
        if current_meeting_context:
            project_data = get_project_by_id(current_meeting_context["project_id"])
            freelancer_data = get_freelancer_by_id(current_meeting_context["freelancer_id"])
            
            if project_data and freelancer_data:
                freelancer_name = f"{freelancer_data['firstName']} {freelancer_data['lastName']}"
                project_name = project_data["projectTitle"]
                
                # Store meeting summary
                # Debug LLM output
                print("🔍 LLM Output:", llm_res["json_output"])
                print("🔍 Summary:", llm_res["json_output"].get("summary", ""))
                print("🔍 Blockers:", llm_res["json_output"].get("blockers", ""))
                
                summary_data = {
                    "freelancer_id": current_meeting_context["freelancer_id"],
                    "project_id": current_meeting_context["project_id"],
                    "freelancer_name": freelancer_name,
                    "project_name": project_name,
                    "summary": llm_res["json_output"].get("summary", ""),
                    "blocker": llm_res["json_output"].get("blockers", "")
                }
                
                store_summary(summary_data)
                print("✅ Meeting summary stored in database")
        
        cleaned = llm_res.get("cleaned_text", raw_txt)
    else:
        cleaned = raw_txt

    # Save to .docx
    doc = Document()
    doc.add_heading(f"Meeting - {datetime.now():%Y-%m-%d %H:%M}", level=0)
    for line in cleaned.splitlines():
        doc.add_paragraph(line)
    doc.save(DOCX_PATH)

    # Upload to GitHub
    # Upload to GitHub Gist (structured JSON-based)
    try:
        from github.uploader import add_meeting
        add_meeting(cleaned)
    except Exception as e:
        print("❌ Gist upload error:", e)

    

    return {"status": "success", "transcript": cleaned}

# ElevenLabs bot automation
def run_bot():
    EMAIL = "aniqaazhar99@gmail.com"
    PASSWORD = "Mazik_Internship@1"
    URL_TO_ADD = "https://gist.githubusercontent.com/SyedAdnanAijaz/1d6b26defc4bb3743567f976d5c0c8fa/raw"

    options = Options()
    options.add_argument('--headless')
    options.add_argument('--disable-gpu')
    options.add_argument('--no-sandbox')
    options.add_argument('--disable-dev-shm-usage')
    options.add_argument('--window-size=1920,1080')
    service = Service(executable_path="chromedriver.exe")

    driver = webdriver.Chrome(service=service, options=options)
    wait = WebDriverWait(driver, 20)

    try:
        driver.get("https://elevenlabs.io/app/sign-in")

        email_input = wait.until(EC.presence_of_element_located((By.CSS_SELECTOR, '[data-testid="sign-in-email-input"]')))
        email_input.send_keys(EMAIL)

        password_input = wait.until(EC.presence_of_element_located((By.CSS_SELECTOR, '[data-testid="sign-in-password-input"]')))
        password_input.send_keys(PASSWORD)

        sign_in_button = wait.until(EC.element_to_be_clickable((By.CSS_SELECTOR, '[data-testid="sign-in-submit-button"]')))
        sign_in_button.click()
        time.sleep(6)
    
        driver.refresh()
        print("🔄 Page refreshed")
        wait.until(EC.presence_of_element_located((By.TAG_NAME, "body")))
        time.sleep(3)

        wait.until(EC.url_contains("/app"))
        conversational_ai = wait.until(EC.element_to_be_clickable((By.CSS_SELECTOR, 'a[aria-label="Conversational AI"]')))
        conversational_ai.click()

        time.sleep(3)
        driver.get("https://elevenlabs.io/app/conversational-ai/agents/agent_4001k1b6f2kefwmtn30k4jcc8gjt")
        wait.until(EC.presence_of_element_located((By.TAG_NAME, "body")))
        time.sleep(3)

        collapse_button = wait.until(EC.element_to_be_clickable((By.CSS_SELECTOR, 'button[aria-label="Collapse"] svg.lucide-trash')))
        driver.execute_script("arguments[0].scrollIntoView({block: 'center', behavior: 'smooth'});", collapse_button)
        time.sleep(1)
        collapse_button.find_element(By.XPATH, "./..").click()

        add_document_button = wait.until(EC.element_to_be_clickable((By.XPATH, '//button[normalize-space()="Add document"]')))
        driver.execute_script("arguments[0].scrollIntoView({block: 'center', behavior: 'smooth'});", add_document_button)
        time.sleep(1)
        add_document_button.click()

        all_buttons = wait.until(EC.presence_of_all_elements_located((By.TAG_NAME, 'button')))
        for button in all_buttons:
            if button.text.strip() == "Add URL":
                driver.execute_script("arguments[0].scrollIntoView({block: 'center', behavior: 'smooth'});", button)
                time.sleep(1)
                button.click()
                break

        url_input = wait.until(EC.presence_of_element_located((By.CSS_SELECTOR, 'input[placeholder="https://example.com"]')))
        url_input.clear()
        url_input.send_keys(URL_TO_ADD)

        final_add_url_button = wait.until(EC.element_to_be_clickable((By.XPATH, '/html/body/form/div[3]/button')))
        driver.execute_script("arguments[0].scrollIntoView({block: 'center', behavior: 'smooth'});", final_add_url_button)
        time.sleep(1)
        final_add_url_button.click()

        time.sleep(10)
        driver.execute_script("window.scrollTo(0, document.body.scrollHeight);")

        unsaved_div = wait.until(EC.presence_of_element_located((By.XPATH, "/html/body/div[1]/div[2]/div[3]/div[3]/div/main/div[2]/form/div[13]/div")))
        save_button = unsaved_div.find_element(By.XPATH, "./div[2]/button[2]")
        driver.execute_script("arguments[0].scrollIntoView({block: 'center', behavior: 'smooth'});", save_button)
        time.sleep(1)
        save_button.click()

        time.sleep(5)

    except Exception as e:
        print("❌ Error:", e)

    finally:
        driver.quit()

# API Endpoints
@app.on_event("startup")
async def startup_event():
    """Initialize database on startup"""
    initialize_database()

@app.get("/")
async def root():
    return {"message": "Meeting Bot API is running"}

@app.get("/cors-test")
async def cors_test():
    """Test endpoint to verify CORS is working"""
    return {"message": "CORS test successful", "timestamp": datetime.now().isoformat()}

@app.post("/start")
def start():
    """Start recording without project context"""
    if is_recording:
        return {"status": "already_recording"}
    start_recording()
    return {"status": "recording_started"}

@app.post("/start-project-meeting")
def start_project_meeting(request: MeetingStartRequest):
    """Start recording with project context"""
    global current_meeting_context
    
    # Validate project and freelancer exist
    project = get_project_by_id(request.project_id)
    freelancer = get_freelancer_by_id(request.freelancer_id)
    
    if not project:
        raise HTTPException(status_code=404, detail="Project not found")
    if not freelancer:
        raise HTTPException(status_code=404, detail="Freelancer not found")
    
    # Set meeting context
    current_meeting_context = {
        "project_id": request.project_id,
        "freelancer_id": request.freelancer_id,
        "project_title": project["projectTitle"],
        "freelancer_name": f"{freelancer['firstName']} {freelancer['lastName']}"
    }
    
    if is_recording:
        return {"status": "already_recording"}
    
    start_recording()
    return {"status": "recording_started", "context": current_meeting_context}

@app.post("/stop")
def stop():
    """Stop recording and generate summary"""
    if not is_recording:
        return {"status": "not_recording"}
    return stop_and_transcribe()

@app.post("/run-elevenlabs-bot")
def run_bot_endpoint():
    """Trigger ElevenLabs bot automation"""
    Thread(target=run_bot).start()
    return {"status": "✅ Bot started in background"}

@app.get("/meeting-summaries")
def get_meeting_summaries(user_id: int, user_role: str, 
                         project_ids: Optional[str] = None, 
                         date_from: Optional[str] = None, 
                         date_to: Optional[str] = None):
    """Get meeting summaries filtered by user and criteria"""
    # Parse project_ids if provided
    project_ids_list = None
    if project_ids:
        project_ids_list = [int(x.strip()) for x in project_ids.split(",")]
    
    result = get_meeting_summaries_by_user(
        user_id, user_role, project_ids_list, date_from, date_to
    )
    return result

@app.post("/store-meeting-summary")
def store_meeting_summary(request: MeetingSummaryRequest):
    """Manually store a meeting summary"""
    summary_data = {
        "freelancer_id": request.freelancer_id,
        "project_id": request.project_id,
        "freelancer_name": request.freelancer_name,
        "project_name": request.project_name,
        "summary": request.summary,
        "blocker": request.blocker
    }
    store_summary(summary_data)
    return {"status": "success", "message": "Meeting summary stored"}

@app.get("/all-meeting-summaries")
def get_all_meeting_summaries():
    """Get all meeting summaries (admin only)"""
    summaries = fetch_all_summaries()
    return {"status": "success", "data": summaries}

@app.get("/project/{project_id}")
def get_project(project_id: int):
    """Get project details"""
    project = get_project_by_id(project_id)
    if not project:
        raise HTTPException(status_code=404, detail="Project not found")
    return project

@app.get("/freelancer/{freelancer_id}")
def get_freelancer(freelancer_id: int):
    """Get freelancer details"""
    freelancer = get_freelancer_by_id(freelancer_id)
    if not freelancer:
        raise HTTPException(status_code=404, detail="Freelancer not found")
    return freelancer

@app.get("/health")
def health_check():
    """Health check endpoint"""
    try:
        # Test database connection
        db_connection = DatabaseConnection()
        connection = db_connection.get_connection()
        connection.close()
        return {"status": "healthy", "database": "connected"}
    except Exception as e:
        return {"status": "unhealthy", "error": str(e)}

@app.get("/summaries")
def get_summaries():
    """Legacy endpoint for backward compatibility"""
    data = fetch_all_summaries()
    return {"status": "success", "data": data}

@app.post("/add-project")
def add_project(payload: ProjectPayload):
    data = load_gist()
 
    for project in data.get("projects", []):
        if project["id"] == payload.id:
            raise HTTPException(
                status_code=400,
                detail=f"❌ Project with ID '{payload.id}' already exists."
            )
 
    new_project = {
        "id": payload.id,
        "projectTitle": payload.projectTitle,
        "description": payload.description,
        "freelancersId": [],
        "freelancersAssigned": []
    }
 
    data.setdefault("projects", []).append(new_project)
    update_gist(data)
 
    return {"message": f"✅ Project '{payload.projectTitle}' added."}
 
# -------------------- API: Add Meeting --------------------
@app.post("/add-meeting")
def add_meeting(payload: MeetingPayload):
    data = load_gist()
 
    timestamp = payload.datetime or datetime.now().strftime("%Y-%m-%d %H:%M")
 
    new_meeting = {
        "datetime": timestamp,
        "transcript": payload.transcript.strip()
    }
 
    data.setdefault("meetings", []).append(new_meeting)
    update_gist(data)
 
    return {"message": f"✅ Meeting added at {timestamp}."}
 
# -------------------- API: Assign Freelancer --------------------
@app.post("/assign-freelancer")
def assign_freelancer(payload: AssignFreelancerPayload):
    data = load_gist()
 
    for project in data.get("projects", []):
        if project["id"] == payload.id:
            # Check if freelancer already assigned
            if payload.freelancerId in project["freelancersId"]:
                return {
                    "message": f"⚠️ Freelancer ID '{payload.freelancerId}' already assigned to project '{payload.id}'."
                }
           
            project["freelancersId"].append(payload.freelancerId)
            project["freelancersAssigned"].append({
                "id": payload.freelancerId,
                "name": payload.freelancerName
            })
 
            update_gist(data)
            return {
                "message": f"✅ Freelancer '{payload.freelancerName}' assigned to project '{payload.id}'."
            }
 
    raise HTTPException(status_code=404, detail=f"❌ Project with ID '{payload.id}' not found.")

if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="0.0.0.0", port=8001)